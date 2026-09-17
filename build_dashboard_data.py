import openpyxl
import json
import sys
import re

def clean_str(val):
    if val is None:
        return ''
    s = str(val).strip()
    s = re.sub(r'[\u200b-\u200d\ufeff]', '', s)
    return s

def sum_valid(lst):
    nums = []
    for x in lst:
        if x is not None and str(x).strip() != '':
            try:
                nums.append(float(x))
            except ValueError:
                pass
    return sum(nums) if nums else 0.0

# Official NT Quality Level Criteria
def get_level_th(pct):
    if pct >= 70.0: return 'ดีมาก'
    elif pct >= 50.0: return 'ดี'
    elif pct >= 29.0: return 'พอใช้'
    else: return 'ปรับปรุง'

def get_level_ma(pct):
    if pct >= 67.0: return 'ดีมาก'
    elif pct >= 44.0: return 'ดี'
    elif pct >= 26.0: return 'พอใช้'
    else: return 'ปรับปรุง'

def get_level_both(pct):
    if pct >= 68.5: return 'ดีมาก'
    elif pct >= 47.0: return 'ดี'
    elif pct >= 27.5: return 'พอใช้'
    else: return 'ปรับปรุง'

def main():
    print("Loading Excel workbook...")
    wb = openpyxl.load_workbook('สรุปผลการทดสอบ NT_ผสาน 17 ศูนย์เครือข่าย_แยกวิชา.xlsx', data_only=True)
    ws_th = wb['ภาษาไทย']
    ws_ma = wb['คณิตศาสตร์']

    students = []
    centers_set = set()
    schools_dict = {}

    for r in range(4, ws_th.max_row + 1):
        c_name = ws_th.cell(r, 1).value
        if not c_name:
            continue
        
        center = clean_str(c_name)
        school = clean_str(ws_th.cell(r, 3).value)
        prefix = clean_str(ws_th.cell(r, 4).value)
        fullname = clean_str(ws_th.cell(r, 5).value)
        stype = clean_str(ws_th.cell(r, 6).value)
        status = clean_str(ws_th.cell(r, 7).value)

        # Standardize stype (ประเภทเด็ก)
        is_special = any(k in stype for k in ['พิเศษ', 'LD', 'บกพร่อง', 'ออทิสติก', 'สมาธิสั้น', 'พิการ'])
        stype_std = 'เด็กพิเศษ' if is_special else 'เด็กปกติ'

        # Standardize status (สถานะการสอบ)
        is_absent = any(k in status for k in ['ขาด', 'ไม่เข้า', 'ย้าย'])
        status_std = 'ขาดสอบ' if is_absent else 'เข้าสอบ'

        # Thai items
        th_choice_items = [ws_th.cell(r, c).value for c in range(8, 34)]
        th_short_items = [ws_th.cell(r, c).value for c in range(36, 39)]
        th_free_item = ws_th.cell(r, 41).value

        # Math items
        ma_choice_items = [ws_ma.cell(r, c).value for c in range(8, 34)]
        ma_short_items = [ws_ma.cell(r, c).value for c in range(36, 39)]
        ma_free_item = ws_ma.cell(r, 41).value

        if status_std == 'เข้าสอบ':
            th_c = round(sum_valid(th_choice_items), 2)
            th_s = round(sum_valid(th_short_items), 2)
            th_f = round(sum_valid([th_free_item]), 2)
            th_tot = round(th_c + th_s + th_f, 2)
            th_pct = round(th_tot, 2)

            ma_c = round(sum_valid(ma_choice_items), 2)
            ma_s = round(sum_valid(ma_short_items), 2)
            ma_f = round(sum_valid([ma_free_item]), 2)
            ma_tot = round(ma_c + ma_s + ma_f, 2)
            ma_pct = round(ma_tot, 2)

            tot_200 = round(th_tot + ma_tot, 2)
            tot_pct = round(tot_200 / 2, 2)
        else:
            th_c = th_s = th_f = th_tot = th_pct = 0.0
            ma_c = ma_s = ma_f = ma_tot = ma_pct = 0.0
            tot_200 = tot_pct = 0.0

        st_obj = {
            'id': len(students) + 1,
            'center': center,
            'school': school,
            'prefix': prefix,
            'name': fullname,
            'type': stype_std,
            'type_orig': stype,
            'status': status_std,
            'status_orig': status,
            'th_c': th_c,
            'th_s': th_s,
            'th_f': th_f,
            'th_tot': th_tot,
            'th_level': get_level_th(th_pct) if status_std == 'เข้าสอบ' else 'ขาดสอบ',
            'ma_c': ma_c,
            'ma_s': ma_s,
            'ma_f': ma_f,
            'ma_tot': ma_tot,
            'ma_level': get_level_ma(ma_pct) if status_std == 'เข้าสอบ' else 'ขาดสอบ',
            'tot_200': tot_200,
            'tot_pct': tot_pct,
            'tot_level': get_level_both(tot_pct) if status_std == 'เข้าสอบ' else 'ขาดสอบ'
        }
        students.append(st_obj)
        centers_set.add(center)

        if school not in schools_dict:
            schools_dict[school] = center

    sorted_centers = sorted(list(centers_set))
    sorted_schools = sorted([{'name': s, 'center': c} for s, c in schools_dict.items()], key=lambda x: (x['center'], x['name']))

    print(f"Total parsed students: {len(students)}")
    print(f"Total unique centers: {len(sorted_centers)}")
    print(f"Total unique schools: {len(sorted_schools)}")

    payload = {
        "centers": sorted_centers,
        "schools": sorted_schools,
        "students": students,
        "criteria": {
            "th": {"excellent": [70.0, 100.0], "good": [50.0, 69.99], "fair": [29.0, 49.99], "needImp": [0.0, 28.99]},
            "ma": {"excellent": [67.0, 100.0], "good": [44.0, 66.99], "fair": [26.0, 43.99], "needImp": [0.0, 25.99]},
            "both": {"excellent": [68.5, 100.0], "good": [47.0, 68.49], "fair": [27.5, 46.99], "needImp": [0.0, 27.49]}
        }
    }

    output_js = f"const NT_DATA = {json.dumps(payload, ensure_ascii=False)};\n"

    with open('nt_data.js', 'w', encoding='utf-8') as f:
        f.write(output_js)

    print("Saved dataset to nt_data.js with official criteria successfully.")

if __name__ == '__main__':
    main()
