import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import re

def clean_str(val):
    if val is None:
        return ''
    s = str(val).strip()
    s = re.sub(r'[\u200b-\u200d\ufeff]', '', s)
    return s

def num_val(v):
    if v is None or v == '':
        return None
    try:
        f = float(v)
        return int(f) if f.is_integer() else round(f, 2)
    except (ValueError, TypeError):
        return None

def main():
    print("Loading source workbook...")
    src_wb = openpyxl.load_workbook('แบบกรอกคะแนนสอบ NTครั้งที่ 1.xlsx', data_only=True)

    all_students = []

    for sheetname in src_wb.sheetnames:
        ws = src_wb[sheetname]
        is_extra = (sheetname == 'กรอกรายชื่อเพิ่มเติม')

        col_school = 3
        col_center = 4 if is_extra else None
        col_prefix = 5 if is_extra else 4
        col_fullname = 6 if is_extra else 5
        col_type = 7 if is_extra else 6
        col_status = 8 if is_extra else 7

        offset = 1 if is_extra else 0

        for r in range(1, ws.max_row + 1):
            name_val = clean_str(ws.cell(r, col_fullname).value)
            if name_val and name_val not in ['ชื่อ - นามสกุล', 'ชื่อ-นามสกุล', 'ชื่อ -นามสกุล']:
                school_val = clean_str(ws.cell(r, col_school).value)
                center_val = clean_str(ws.cell(r, col_center).value) if is_extra else sheetname
                
                # Map typo 'ชุมแพชนูปถัมภ์' center to 'ชุมแพ'
                if center_val == 'ชุมแพชนูปถัมภ์':
                    center_val = 'ชุมแพ'

                prefix_val = clean_str(ws.cell(r, col_prefix).value)
                type_val = clean_str(ws.cell(r, col_type).value)
                status_val = clean_str(ws.cell(r, col_status).value)

                # Thai scores (Cols 8 to 45)
                th_q_choice = [num_val(ws.cell(r, c + offset).value) for c in range(8, 34)]
                th_tot_choice = num_val(ws.cell(r, 34 + offset).value)
                th_pct_choice = num_val(ws.cell(r, 35 + offset).value)
                th_q_short = [num_val(ws.cell(r, c + offset).value) for c in range(36, 39)]
                th_tot_short = num_val(ws.cell(r, 39 + offset).value)
                th_pct_short = num_val(ws.cell(r, 40 + offset).value)
                th_q_free = num_val(ws.cell(r, 41 + offset).value)
                th_tot_free = num_val(ws.cell(r, 42 + offset).value)
                th_pct_free = num_val(ws.cell(r, 43 + offset).value)
                th_tot_100 = num_val(ws.cell(r, 44 + offset).value)
                th_pct_100 = num_val(ws.cell(r, 45 + offset).value)

                # Math scores (Cols 46 to 83)
                ma_q_choice = [num_val(ws.cell(r, c + offset).value) for c in range(46, 72)]
                ma_tot_choice = num_val(ws.cell(r, 72 + offset).value)
                ma_pct_choice = num_val(ws.cell(r, 73 + offset).value)
                ma_q_short = [num_val(ws.cell(r, c + offset).value) for c in range(74, 77)]
                ma_tot_short = num_val(ws.cell(r, 77 + offset).value)
                ma_pct_short = num_val(ws.cell(r, 78 + offset).value)
                ma_q_free = num_val(ws.cell(r, 79 + offset).value)
                ma_tot_free = num_val(ws.cell(r, 80 + offset).value)
                ma_pct_free = num_val(ws.cell(r, 81 + offset).value)
                ma_tot_100 = num_val(ws.cell(r, 82 + offset).value)
                ma_pct_100 = num_val(ws.cell(r, 83 + offset).value)

                tot_2subj = num_val(ws.cell(r, 84 + offset).value)
                pct_2subj = num_val(ws.cell(r, 85 + offset).value)

                rec = {
                    'center': center_val,
                    'school': school_val,
                    'prefix': prefix_val,
                    'fullname': name_val,
                    'type': type_val,
                    'status': status_val,
                    # Thai
                    'th_q_choice': th_q_choice,
                    'th_tot_choice': th_tot_choice,
                    'th_pct_choice': th_pct_choice,
                    'th_q_short': th_q_short,
                    'th_tot_short': th_tot_short,
                    'th_pct_short': th_pct_short,
                    'th_q_free': th_q_free,
                    'th_tot_free': th_tot_free,
                    'th_pct_free': th_pct_free,
                    'th_tot_100': th_tot_100,
                    'th_pct_100': th_pct_100,
                    # Math
                    'ma_q_choice': ma_q_choice,
                    'ma_tot_choice': ma_tot_choice,
                    'ma_pct_choice': ma_pct_choice,
                    'ma_q_short': ma_q_short,
                    'ma_tot_short': ma_tot_short,
                    'ma_pct_short': ma_pct_short,
                    'ma_q_free': ma_q_free,
                    'ma_tot_free': ma_tot_free,
                    'ma_pct_free': ma_pct_free,
                    'ma_tot_100': ma_tot_100,
                    'ma_pct_100': ma_pct_100,
                    # Both
                    'tot_2subj': tot_2subj,
                    'pct_2subj': pct_2subj,
                }
                all_students.append(rec)

    print(f"Extracted {len(all_students)} total student records across 17 centers.")

    # Sort students by Center, School, Fullname
    all_students.sort(key=lambda x: (x['center'], x['school'], x['fullname']))

    out_wb = openpyxl.Workbook()
    out_wb.remove(out_wb.active)

    # Styling definitions
    font_family = "Segoe UI"
    f_title = Font(name=font_family, size=14, bold=True, color="FFFFFF")
    f_group = Font(name=font_family, size=11, bold=True, color="FFFFFF")
    f_col = Font(name=font_family, size=10, bold=True, color="1B2631")
    f_data = Font(name=font_family, size=10, color="000000")
    f_bold = Font(name=font_family, size=10, bold=True, color="000000")

    fill_title = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    
    fill_grp_meta = PatternFill(start_color="34495E", end_color="34495E", fill_type="solid")
    fill_grp_th = PatternFill(start_color="1E8449", end_color="1E8449", fill_type="solid")
    fill_grp_ma = PatternFill(start_color="2980B9", end_color="2980B9", fill_type="solid")
    fill_grp_both = PatternFill(start_color="6C3483", end_color="6C3483", fill_type="solid")

    fill_hdr_meta = PatternFill(start_color="EAECEE", end_color="EAECEE", fill_type="solid")
    fill_hdr_th = PatternFill(start_color="D4EFDF", end_color="D4EFDF", fill_type="solid")
    fill_hdr_ma = PatternFill(start_color="D6EAF8", end_color="D6EAF8", fill_type="solid")
    fill_hdr_both = PatternFill(start_color="EBDEF0", end_color="EBDEF0", fill_type="solid")

    fill_zebra = PatternFill(start_color="F8F9F9", end_color="F8F9F9", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='D5D8DC'),
        right=Side(style='thin', color='D5D8DC'),
        top=Side(style='thin', color='D5D8DC'),
        bottom=Side(style='thin', color='D5D8DC')
    )

    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_left = Alignment(horizontal='left', vertical='center')
    align_right = Alignment(horizontal='right', vertical='center')

    def style_header_cell(cell, font=f_col, fill=fill_hdr_meta, align=align_center):
        cell.font = font
        cell.fill = fill
        cell.alignment = align
        cell.border = thin_border

    # ==========================================
    # SHEET 1: ภาษาไทย
    # ==========================================
    ws_th = out_wb.create_sheet(title="ภาษาไทย")
    ws_th.views.sheetView[0].showGridLines = True

    # Title Banner Row 1
    ws_th.merge_cells("A1:AS1")
    t_cell = ws_th.cell(1, 1, "ผลการทดสอบความสามารถด้านภาษาไทย (NT) ชั้นประถมศึกษาปีที่ 3 รวม 17 ศูนย์เครือข่าย")
    t_cell.font = f_title
    t_cell.fill = fill_title
    t_cell.alignment = align_center
    ws_th.row_dimensions[1].height = 35

    # Group Headers Row 2
    ws_th.merge_cells("A2:G2")
    ws_th.cell(2, 1, "ข้อมูลนักเรียน").font = f_group
    for c in range(1, 8): ws_th.cell(2, c).fill = fill_grp_meta; ws_th.cell(2, c).alignment = align_center

    ws_th.merge_cells("H2:AI2")
    ws_th.cell(2, 8, "แบบเลือกตอบ (78 คะแนน)").font = f_group
    for c in range(8, 36): ws_th.cell(2, c).fill = fill_grp_th; ws_th.cell(2, c).alignment = align_center

    ws_th.merge_cells("AJ2:AN2")
    ws_th.cell(2, 36, "เขียนตอบสั้น (15 คะแนน)").font = f_group
    for c in range(36, 41): ws_th.cell(2, c).fill = fill_grp_th; ws_th.cell(2, c).alignment = align_center

    ws_th.merge_cells("AO2:AQ2")
    ws_th.cell(2, 41, "เขียนตอบอิสระ (7 คะแนน)").font = f_group
    for c in range(41, 44): ws_th.cell(2, c).fill = fill_grp_th; ws_th.cell(2, c).alignment = align_center

    ws_th.merge_cells("AR2:AS2")
    ws_th.cell(2, 44, "สรุปผลวิชาภาษาไทย (100 คะแนน)").font = f_group
    for c in range(44, 46): ws_th.cell(2, c).fill = fill_grp_th; ws_th.cell(2, c).alignment = align_center

    ws_th.row_dimensions[2].height = 25

    # Column Headers Row 3
    th_headers = [
        "ศูนย์เครือข่าย", "ลำดับที่", "โรงเรียน", "คำนำหน้า", "ชื่อ - นามสกุล", "ประเภทเด็ก*", "สถานะ การสอบ**"
    ] + [f"ข้อที่ {i}" for i in range(1, 27)] + ["รวม (78 คะแนน)", "ร้อยละ"] \
      + [f"ข้อที่ {i}" for i in range(27, 30)] + ["รวม (15 คะแนน)", "ร้อยละ"] \
      + ["ข้อที่ 30", "รวม (7 คะแนน)", "ร้อยละ"] \
      + ["รวม 100 คะแนน", "ร้อยละ"]

    for col_idx, htext in enumerate(th_headers, 1):
        c_cell = ws_th.cell(3, col_idx, htext)
        fill_c = fill_hdr_meta if col_idx <= 7 else fill_hdr_th
        style_header_cell(c_cell, font=f_col, fill=fill_c, align=align_center)
    ws_th.row_dimensions[3].height = 28

    # Populate Data Rows
    row_no = 4
    for idx, st in enumerate(all_students, 1):
        ws_th.cell(row_no, 1, st['center'])
        ws_th.cell(row_no, 2, idx)
        ws_th.cell(row_no, 3, st['school'])
        ws_th.cell(row_no, 4, st['prefix'])
        ws_th.cell(row_no, 5, st['fullname'])
        ws_th.cell(row_no, 6, st['type'])
        ws_th.cell(row_no, 7, st['status'])

        # Choice Q1-26
        for i, val in enumerate(st['th_q_choice']):
            ws_th.cell(row_no, 8 + i, val)
        ws_th.cell(row_no, 34, f"=SUM(H{row_no}:AG{row_no})")
        ws_th.cell(row_no, 35, f"=AH{row_no}/78*100")

        # Short Q27-29
        for i, val in enumerate(st['th_q_short']):
            ws_th.cell(row_no, 36 + i, val)
        ws_th.cell(row_no, 39, f"=SUM(AJ{row_no}:AL{row_no})")
        ws_th.cell(row_no, 40, f"=AM{row_no}/15*100")

        # Free Q30
        ws_th.cell(row_no, 41, st['th_q_free'])
        ws_th.cell(row_no, 42, f"=AO{row_no}")
        ws_th.cell(row_no, 43, f"=AP{row_no}/7*100")

        # Total 100
        ws_th.cell(row_no, 44, f"=AH{row_no}+AM{row_no}+AP{row_no}")
        ws_th.cell(row_no, 45, f"=AR{row_no}")

        # Formatting data cells
        fill_r = fill_zebra if row_no % 2 == 0 else PatternFill(fill_type=None)
        for col_idx in range(1, 46):
            cell = ws_th.cell(row_no, col_idx)
            cell.font = f_data
            cell.border = thin_border
            if fill_r.fill_type: cell.fill = fill_r
            
            if col_idx in [1, 3, 4, 5]:
                cell.alignment = align_left
            elif col_idx in [2, 6, 7]:
                cell.alignment = align_center
            else:
                cell.alignment = align_right
                if col_idx in [35, 40, 43, 45]:
                    cell.number_format = '0.00'
                elif col_idx in [34, 39, 42, 44]:
                    cell.number_format = '0.0'
                else:
                    cell.number_format = '0'

        row_no += 1

    ws_th.freeze_panes = 'F4'

    # ==========================================
    # SHEET 2: คณิตศาสตร์
    # ==========================================
    ws_ma = out_wb.create_sheet(title="คณิตศาสตร์")
    ws_ma.views.sheetView[0].showGridLines = True

    # Title Banner Row 1
    ws_ma.merge_cells("A1:AS1")
    t_cell = ws_ma.cell(1, 1, "ผลการทดสอบความสามารถด้านคณิตศาสตร์ (NT) ชั้นประถมศึกษาปีที่ 3 รวม 17 ศูนย์เครือข่าย")
    t_cell.font = f_title
    t_cell.fill = fill_title
    t_cell.alignment = align_center
    ws_ma.row_dimensions[1].height = 35

    # Group Headers Row 2
    ws_ma.merge_cells("A2:G2")
    ws_ma.cell(2, 1, "ข้อมูลนักเรียน").font = f_group
    for c in range(1, 8): ws_ma.cell(2, c).fill = fill_grp_meta; ws_ma.cell(2, c).alignment = align_center

    ws_ma.merge_cells("H2:AI2")
    ws_ma.cell(2, 8, "แบบเลือกตอบ (78 คะแนน)").font = f_group
    for c in range(8, 36): ws_ma.cell(2, c).fill = fill_grp_ma; ws_ma.cell(2, c).alignment = align_center

    ws_ma.merge_cells("AJ2:AN2")
    ws_ma.cell(2, 36, "เขียนตอบสั้น (12 คะแนน)").font = f_group
    for c in range(36, 41): ws_ma.cell(2, c).fill = fill_grp_ma; ws_ma.cell(2, c).alignment = align_center

    ws_ma.merge_cells("AO2:AQ2")
    ws_ma.cell(2, 41, "เขียนตอบอิสระ (10 คะแนน)").font = f_group
    for c in range(41, 44): ws_ma.cell(2, c).fill = fill_grp_ma; ws_ma.cell(2, c).alignment = align_center

    ws_ma.merge_cells("AR2:AS2")
    ws_ma.cell(2, 44, "สรุปผลวิชาคณิตศาสตร์ (100 คะแนน)").font = f_group
    for c in range(44, 46): ws_ma.cell(2, c).fill = fill_grp_ma; ws_ma.cell(2, c).alignment = align_center

    ws_ma.row_dimensions[2].height = 25

    # Column Headers Row 3
    ma_headers = [
        "ศูนย์เครือข่าย", "ลำดับที่", "โรงเรียน", "คำนำหน้า", "ชื่อ - นามสกุล", "ประเภทเด็ก*", "สถานะ การสอบ**"
    ] + [f"ข้อที่ {i}" for i in range(1, 27)] + ["รวม (78 คะแนน)", "ร้อยละ"] \
      + [f"ข้อที่ {i}" for i in range(27, 30)] + ["รวม (12 คะแนน)", "ร้อยละ"] \
      + ["ข้อที่ 30", "รวม (10 คะแนน)", "ร้อยละ"] \
      + ["รวม 100 คะแนน", "ร้อยละ"]

    for col_idx, htext in enumerate(ma_headers, 1):
        c_cell = ws_ma.cell(3, col_idx, htext)
        fill_c = fill_hdr_meta if col_idx <= 7 else fill_hdr_ma
        style_header_cell(c_cell, font=f_col, fill=fill_c, align=align_center)
    ws_ma.row_dimensions[3].height = 28

    # Populate Data Rows
    row_no = 4
    for idx, st in enumerate(all_students, 1):
        ws_ma.cell(row_no, 1, st['center'])
        ws_ma.cell(row_no, 2, idx)
        ws_ma.cell(row_no, 3, st['school'])
        ws_ma.cell(row_no, 4, st['prefix'])
        ws_ma.cell(row_no, 5, st['fullname'])
        ws_ma.cell(row_no, 6, st['type'])
        ws_ma.cell(row_no, 7, st['status'])

        # Choice Q1-26
        for i, val in enumerate(st['ma_q_choice']):
            ws_ma.cell(row_no, 8 + i, val)
        ws_ma.cell(row_no, 34, f"=SUM(H{row_no}:AG{row_no})")
        ws_ma.cell(row_no, 35, f"=AH{row_no}/78*100")

        # Short Q27-29
        for i, val in enumerate(st['ma_q_short']):
            ws_ma.cell(row_no, 36 + i, val)
        ws_ma.cell(row_no, 39, f"=SUM(AJ{row_no}:AL{row_no})")
        ws_ma.cell(row_no, 40, f"=AM{row_no}/12*100")

        # Free Q30
        ws_ma.cell(row_no, 41, st['ma_q_free'])
        ws_ma.cell(row_no, 42, f"=AO{row_no}")
        ws_ma.cell(row_no, 43, f"=AP{row_no}/10*100")

        # Total 100
        ws_ma.cell(row_no, 44, f"=AH{row_no}+AM{row_no}+AP{row_no}")
        ws_ma.cell(row_no, 45, f"=AR{row_no}")

        # Formatting data cells
        fill_r = fill_zebra if row_no % 2 == 0 else PatternFill(fill_type=None)
        for col_idx in range(1, 46):
            cell = ws_ma.cell(row_no, col_idx)
            cell.font = f_data
            cell.border = thin_border
            if fill_r.fill_type: cell.fill = fill_r
            
            if col_idx in [1, 3, 4, 5]:
                cell.alignment = align_left
            elif col_idx in [2, 6, 7]:
                cell.alignment = align_center
            else:
                cell.alignment = align_right
                if col_idx in [35, 40, 43, 45]:
                    cell.number_format = '0.00'
                elif col_idx in [34, 39, 42, 44]:
                    cell.number_format = '0.0'
                else:
                    cell.number_format = '0'

        row_no += 1

    ws_ma.freeze_panes = 'F4'

    # ==========================================
    # SHEET 3: ภาพรวม 2 วิชา
    # ==========================================
    ws_both = out_wb.create_sheet(title="ภาพรวม 2 วิชา")
    ws_both.views.sheetView[0].showGridLines = True

    # Title Banner Row 1
    ws_both.merge_cells("A1:S1")
    t_cell = ws_both.cell(1, 1, "สรุปภาพรวมผลการทดสอบ NT (ภาษาไทย และ คณิตศาสตร์) รวม 17 ศูนย์เครือข่าย")
    t_cell.font = f_title
    t_cell.fill = fill_title
    t_cell.alignment = align_center
    ws_both.row_dimensions[1].height = 35

    # Group Headers Row 2
    ws_both.merge_cells("A2:G2")
    ws_both.cell(2, 1, "ข้อมูลนักเรียน").font = f_group
    for c in range(1, 8): ws_both.cell(2, c).fill = fill_grp_meta; ws_both.cell(2, c).alignment = align_center

    ws_both.merge_cells("H2:L2")
    ws_both.cell(2, 8, "วิชาภาษาไทย (100 คะแนน)").font = f_group
    for c in range(8, 13): ws_both.cell(2, c).fill = fill_grp_th; ws_both.cell(2, c).alignment = align_center

    ws_both.merge_cells("M2:Q2")
    ws_both.cell(2, 13, "วิชาคณิตศาสตร์ (100 คะแนน)").font = f_group
    for c in range(13, 18): ws_both.cell(2, c).fill = fill_grp_ma; ws_both.cell(2, c).alignment = align_center

    ws_both.merge_cells("R2:S2")
    ws_both.cell(2, 18, "รวม 2 วิชา (200 คะแนน)").font = f_group
    for c in range(18, 20): ws_both.cell(2, c).fill = fill_grp_both; ws_both.cell(2, c).alignment = align_center

    ws_both.row_dimensions[2].height = 25

    # Column Headers Row 3
    both_headers = [
        "ศูนย์เครือข่าย", "ลำดับที่", "โรงเรียน", "คำนำหน้า", "ชื่อ - นามสกุล", "ประเภทเด็ก*", "สถานะ การสอบ**",
        "เลือกตอบ (78)", "ตอบสั้น (15)", "ตอบอิสระ (7)", "รวมภาษาไทย (100)", "ร้อยละ ภาษาไทย",
        "เลือกตอบ (78)", "ตอบสั้น (12)", "ตอบอิสระ (10)", "รวมคณิตศาสตร์ (100)", "ร้อยละ คณิตศาสตร์",
        "รวม 2 วิชา (200)", "ร้อยละ รวม 2 วิชา"
    ]

    for col_idx, htext in enumerate(both_headers, 1):
        c_cell = ws_both.cell(3, col_idx, htext)
        fill_c = fill_hdr_meta if col_idx <= 7 else (fill_hdr_th if col_idx <= 12 else (fill_hdr_ma if col_idx <= 17 else fill_hdr_both))
        style_header_cell(c_cell, font=f_col, fill=fill_c, align=align_center)
    ws_both.row_dimensions[3].height = 28

    # Populate Data Rows
    row_no = 4
    for idx, st in enumerate(all_students, 1):
        ws_both.cell(row_no, 1, st['center'])
        ws_both.cell(row_no, 2, idx)
        ws_both.cell(row_no, 3, st['school'])
        ws_both.cell(row_no, 4, st['prefix'])
        ws_both.cell(row_no, 5, st['fullname'])
        ws_both.cell(row_no, 6, st['type'])
        ws_both.cell(row_no, 7, st['status'])

        # Formulas linking to Thai and Math sheets
        ws_both.cell(row_no, 8, f"=ภาษาไทย!AH{row_no}")
        ws_both.cell(row_no, 9, f"=ภาษาไทย!AM{row_no}")
        ws_both.cell(row_no, 10, f"=ภาษาไทย!AP{row_no}")
        ws_both.cell(row_no, 11, f"=ภาษาไทย!AR{row_no}")
        ws_both.cell(row_no, 12, f"=ภาษาไทย!AS{row_no}")

        ws_both.cell(row_no, 13, f"=คณิตศาสตร์!AH{row_no}")
        ws_both.cell(row_no, 14, f"=คณิตศาสตร์!AM{row_no}")
        ws_both.cell(row_no, 15, f"=คณิตศาสตร์!AP{row_no}")
        ws_both.cell(row_no, 16, f"=คณิตศาสตร์!AR{row_no}")
        ws_both.cell(row_no, 17, f"=คณิตศาสตร์!AS{row_no}")

        ws_both.cell(row_no, 18, f"=K{row_no}+P{row_no}")
        ws_both.cell(row_no, 19, f"=R{row_no}/200*100")

        # Formatting data cells
        fill_r = fill_zebra if row_no % 2 == 0 else PatternFill(fill_type=None)
        for col_idx in range(1, 20):
            cell = ws_both.cell(row_no, col_idx)
            cell.font = f_data
            cell.border = thin_border
            if fill_r.fill_type: cell.fill = fill_r
            
            if col_idx in [1, 3, 4, 5]:
                cell.alignment = align_left
            elif col_idx in [2, 6, 7]:
                cell.alignment = align_center
            else:
                cell.alignment = align_right
                if col_idx in [12, 17, 19]:
                    cell.number_format = '0.00'
                else:
                    cell.number_format = '0.0'

        row_no += 1

    ws_both.freeze_panes = 'F4'

    # ==========================================
    # SHEET 4: สรุปตามศูนย์เครือข่าย
    # ==========================================
    ws_sum = out_wb.create_sheet(title="สรุปตามศูนย์เครือข่าย")
    ws_sum.views.sheetView[0].showGridLines = True

    # Title Banner Row 1
    ws_sum.merge_cells("A1:K1")
    t_cell = ws_sum.cell(1, 1, "สรุปผลการทดสอบ NT รายศูนย์เครือข่าย (17 ศูนย์เครือข่าย)")
    t_cell.font = f_title
    t_cell.fill = fill_title
    t_cell.alignment = align_center
    ws_sum.row_dimensions[1].height = 35

    # Group Headers Row 2
    ws_sum.merge_cells("A2:E2")
    ws_sum.cell(2, 1, "ข้อมูลศูนย์เครือข่าย").font = f_group
    for c in range(1, 6): ws_sum.cell(2, c).fill = fill_grp_meta; ws_sum.cell(2, c).alignment = align_center

    ws_sum.merge_cells("F2:G2")
    ws_sum.cell(2, 6, "เฉลี่ย วิชาภาษาไทย").font = f_group
    for c in range(6, 8): ws_sum.cell(2, c).fill = fill_grp_th; ws_sum.cell(2, c).alignment = align_center

    ws_sum.merge_cells("H2:I2")
    ws_sum.cell(2, 8, "เฉลี่ย วิชาคณิตศาสตร์").font = f_group
    for c in range(8, 10): ws_sum.cell(2, c).fill = fill_grp_ma; ws_sum.cell(2, c).alignment = align_center

    ws_sum.merge_cells("J2:K2")
    ws_sum.cell(2, 10, "เฉลี่ย รวม 2 วิชา").font = f_group
    for c in range(10, 12): ws_sum.cell(2, c).fill = fill_grp_both; ws_sum.cell(2, c).alignment = align_center

    ws_sum.row_dimensions[2].height = 25

    # Column Headers Row 3
    sum_headers = [
        "ที่", "ศูนย์เครือข่าย", "จำนวนโรงเรียน", "จำนวนนักเรียน (คน)", "ผู้เข้าสอบ (คน)",
        "คะแนนเฉลี่ย (100)", "ร้อยละเฉลี่ย",
        "คะแนนเฉลี่ย (100)", "ร้อยละเฉลี่ย",
        "คะแนนเฉลี่ย (200)", "ร้อยละเฉลี่ย"
    ]

    for col_idx, htext in enumerate(sum_headers, 1):
        c_cell = ws_sum.cell(3, col_idx, htext)
        fill_c = fill_hdr_meta if col_idx <= 5 else (fill_hdr_th if col_idx <= 7 else (fill_hdr_ma if col_idx <= 9 else fill_hdr_both))
        style_header_cell(c_cell, font=f_col, fill=fill_c, align=align_center)
    ws_sum.row_dimensions[3].height = 28

    # Get distinct 17 centers in alphabetical order
    centers_list = sorted(list(set(st['center'] for st in all_students)))

    # Calculate pre-computed unique school counts per center for maximum Excel compatibility
    school_counts_per_center = {}
    for c_name in centers_list:
        schs = set(st['school'] for st in all_students if st['center'] == c_name)
        school_counts_per_center[c_name] = len(schs)

    row_no = 4
    for idx, c_name in enumerate(centers_list, 1):
        ws_sum.cell(row_no, 1, idx)
        ws_sum.cell(row_no, 2, c_name)
        
        # Pre-computed school count
        ws_sum.cell(row_no, 3, school_counts_per_center[c_name])
        
        # Standard formulas compatible with all Excel versions
        ws_sum.cell(row_no, 4, f'=COUNTIF(\'ภาพรวม 2 วิชา\'!$A$4:$A$3086, "{c_name}")')
        ws_sum.cell(row_no, 5, f'=COUNTIFS(\'ภาพรวม 2 วิชา\'!$A$4:$A$3086, "{c_name}", \'ภาพรวม 2 วิชา\'!$G$4:$G$3086, "*เข้าสอบ*")')
        
        ws_sum.cell(row_no, 6, f'=AVERAGEIFS(\'ภาพรวม 2 วิชา\'!$K$4:$K$3086, \'ภาพรวม 2 วิชา\'!$A$4:$A$3086, "{c_name}", \'ภาพรวม 2 วิชา\'!$G$4:$G$3086, "*เข้าสอบ*")')
        ws_sum.cell(row_no, 7, f'=F{row_no}')
        
        ws_sum.cell(row_no, 8, f'=AVERAGEIFS(\'ภาพรวม 2 วิชา\'!$P$4:$P$3086, \'ภาพรวม 2 วิชา\'!$A$4:$A$3086, "{c_name}", \'ภาพรวม 2 วิชา\'!$G$4:$G$3086, "*เข้าสอบ*")')
        ws_sum.cell(row_no, 9, f'=H{row_no}')
        
        ws_sum.cell(row_no, 10, f'=AVERAGEIFS(\'ภาพรวม 2 วิชา\'!$R$4:$R$3086, \'ภาพรวม 2 วิชา\'!$A$4:$A$3086, "{c_name}", \'ภาพรวม 2 วิชา\'!$G$4:$G$3086, "*เข้าสอบ*")')
        ws_sum.cell(row_no, 11, f'=J{row_no}/2')

        fill_r = fill_zebra if row_no % 2 == 0 else PatternFill(fill_type=None)
        for col_idx in range(1, 12):
            cell = ws_sum.cell(row_no, col_idx)
            cell.font = f_data
            cell.border = thin_border
            if fill_r.fill_type: cell.fill = fill_r
            
            if col_idx == 2:
                cell.alignment = align_left
            elif col_idx in [1, 3, 4, 5]:
                cell.alignment = align_center
                cell.number_format = '#,##0'
            else:
                cell.alignment = align_right
                cell.number_format = '0.00'

        row_no += 1

    # Total Summary Row at bottom
    ws_sum.cell(row_no, 1, "")
    tot_c = ws_sum.cell(row_no, 2, "รวมทั้งสิ้น / เฉลี่ยภาพรวม")
    tot_c.font = f_bold
    tot_c.alignment = align_center
    
    total_unique_schools = len(set(st['school'] for st in all_students))
    ws_sum.cell(row_no, 3, total_unique_schools)
    ws_sum.cell(row_no, 4, f'=SUM(D4:D{row_no-1})')
    ws_sum.cell(row_no, 5, f'=SUM(E4:E{row_no-1})')
    
    ws_sum.cell(row_no, 6, f'=AVERAGEIFS(\'ภาพรวม 2 วิชา\'!$K$4:$K$3086, \'ภาพรวม 2 วิชา\'!$G$4:$G$3086, "*เข้าสอบ*")')
    ws_sum.cell(row_no, 7, f'=F{row_no}')
    
    ws_sum.cell(row_no, 8, f'=AVERAGEIFS(\'ภาพรวม 2 วิชา\'!$P$4:$P$3086, \'ภาพรวม 2 วิชา\'!$G$4:$G$3086, "*เข้าสอบ*")')
    ws_sum.cell(row_no, 9, f'=H{row_no}')
    
    ws_sum.cell(row_no, 10, f'=AVERAGEIFS(\'ภาพรวม 2 วิชา\'!$R$4:$R$3086, \'ภาพรวม 2 วิชา\'!$G$4:$G$3086, "*เข้าสอบ*")')
    ws_sum.cell(row_no, 11, f'=J{row_no}/2')

    fill_tot = PatternFill(start_color="EAEDED", end_color="EAEDED", fill_type="solid")
    for col_idx in range(1, 12):
        cell = ws_sum.cell(row_no, col_idx)
        cell.font = f_bold
        cell.fill = fill_tot
        cell.border = thin_border
        if col_idx in [3, 4, 5]:
            cell.alignment = align_center
            cell.number_format = '#,##0'
        elif col_idx >= 6:
            cell.alignment = align_right
            cell.number_format = '0.00'

    ws_sum.row_dimensions[row_no].height = 25

    # Adjust Column Widths across all sheets
    for ws in [ws_th, ws_ma, ws_both, ws_sum]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row in [1, 2]: continue # ignore title banner
                val = str(cell.value or '')
                length = sum(2 if ord(ch) > 255 else 1 for ch in val)
                if length > max_len: max_len = length
            ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

    # Set specific column widths for metadata
    ws_th.column_dimensions['A'].width = 24
    ws_th.column_dimensions['C'].width = 30
    ws_th.column_dimensions['E'].width = 26

    ws_ma.column_dimensions['A'].width = 24
    ws_ma.column_dimensions['C'].width = 30
    ws_ma.column_dimensions['E'].width = 26

    ws_both.column_dimensions['A'].width = 24
    ws_both.column_dimensions['C'].width = 30
    ws_both.column_dimensions['E'].width = 26

    ws_sum.column_dimensions['A'].width = 8
    ws_sum.column_dimensions['B'].width = 28
    ws_sum.column_dimensions['C'].width = 16
    ws_sum.column_dimensions['D'].width = 20
    ws_sum.column_dimensions['E'].width = 18

    output_filename = 'สรุปผลการทดสอบ NT_ผสาน 17 ศูนย์เครือข่าย_แยกวิชา.xlsx'
    out_wb.save(output_filename)
    print(f"Successfully saved output file to: {output_filename}")

if __name__ == '__main__':
    main()
